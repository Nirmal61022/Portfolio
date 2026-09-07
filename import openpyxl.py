import openpyxl
def audit_cell_map():
    wb = openpyxl.load_workbook(r"C:\Users\nirma\.gemini\antigravity\scratch\emlt_lsm_workbook\EMLT_LSM_Engineering_Calculation_Workbook.xlsx", data_only=False)
    
    print("--- INPUTS SHEET AUDIT ---")
    ws_inp = wb["INPUTS"]
    for r in range(4, ws_inp.max_row + 1):
        c1 = ws_inp.cell(row=r, column=1).value
        c2 = ws_inp.cell(row=r, column=2).value
        c3 = ws_inp.cell(row=r, column=3).value
        c4 = ws_inp.cell(row=r, column=4).value
        if c2:
            print(f"Row {r:2d} | Sym: {c3:12s} | Name: {c2:35s} | Val/Formula: {c4}")
    print("\n--- MECHANICAL_CALCS SHEET AUDIT ---")
    ws_mech = wb["MECHANICAL_CALCS"]
    for r in range(4, 23):
        c1 = ws_mech.cell(row=r, column=1).value
        c2 = ws_mech.cell(row=r, column=2).value
        c3 = ws_mech.cell(row=r, column=3).value
        if c1:
            print(f"Row {r:2d} | Sym: {c2:12s} | Name: {c1:35s} | Val/Formula: {c3}")
    print("\n--- MAGNETIC_CALCS SHEET AUDIT ---")
    ws_mag = wb["MAGNETIC_CALCS"]
    for r in range(4, 13):
        c1 = ws_mag.cell(row=r, column=1).value
        c2 = ws_mag.cell(row=r, column=2).value
        c3 = ws_mag.cell(row=r, column=3).value
        if c1:
            print(f"Row {r:2d} | Sym: {str(c2):12s} | Name: {c1:45s} | Val/Formula: {c3}")
    print("\n--- COIL_DESIGN SHEET AUDIT ---")
    ws_coil = wb["COIL_DESIGN"]
    for r in range(4, 17):
        c1 = ws_coil.cell(row=r, column=1).value
        c2 = ws_coil.cell(row=r, column=2).value
        c3 = ws_coil.cell(row=r, column=3).value